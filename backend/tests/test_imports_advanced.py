"""Advanced imports: XLSX, saved mappings, error export, async commit.

The claims under test are the ones a seller's data depends on:

* a spreadsheet reads exactly like the CSV of the same sheet, so money and
  phone parsing cannot drift between the two routes;
* a failed row comes back as a file the seller can fix and re-upload;
* a commit is idempotent and resumable, so a retry never doubles an order;
* a large import is handed to the worker inside the request's transaction, so a
  rolled-back request leaves no job behind.

The spreadsheet and export tests are pure; the rest run against the real
database through the service, because the HTTP fixture the V1 suites use is
broken in this environment for unrelated reasons.
"""

from __future__ import annotations

import datetime as dt
import io
import uuid

import pytest
import sqlalchemy as sa
from openpyxl import Workbook

from app.common.outbox import OutboxEvent, OutboxTopic
from app.common.uploads import DetectedFormat, validate_upload
from app.core.clock import utc_now
from app.core.context import RequestContext, set_context
from app.core.errors import ConflictError, ValidationError
from app.imports.models import (
    ImportBatch,
    ImportMapping,
    ImportRow,
    ImportRowStatus,
    ImportStatus,
    ImportTemplate,
)
from app.imports.service import (
    ASYNC_COMMIT_THRESHOLD_ROWS,
    build_import_service,
)
from app.imports.spreadsheet import cell_to_text, read_xlsx
from app.tenants.models import Tenant


def _workbook(rows: list[list[object]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def service(db, settings):
    return build_import_service(db, settings)


async def _shop(db) -> Tenant:
    tenant = Tenant(name="Import Shop")
    db.add(tenant)
    await db.flush()
    set_context(RequestContext(trace_id="test", tenant_id=tenant.id))
    return tenant


# ------------------------------------------------------------- spreadsheet --


def test_a_whole_number_does_not_gain_decimals() -> None:
    """The reason V1 refused spreadsheets.

    Excel stores 1050 as a float. Handing 1050.0 to the money parser is how an
    amount a seller typed becomes a different number.
    """
    assert cell_to_text(1050.0) == "1050"
    assert cell_to_text(1050) == "1050"
    assert cell_to_text(1050.5) == "1050.5"
    # Trailing zeros a float repr leaves behind are dropped, not carried.
    assert cell_to_text(1050.50) == "1050.5"


def test_a_date_cell_becomes_a_date_not_a_serial_number() -> None:
    assert cell_to_text(dt.datetime(2026, 9, 18)) == "2026-09-18"
    assert cell_to_text(dt.datetime(2026, 9, 18, 14, 30)) == "2026-09-18 14:30:00"
    assert cell_to_text(dt.date(2026, 9, 18)) == "2026-09-18"


def test_an_empty_or_unreadable_cell_becomes_empty_never_a_guess() -> None:
    """An uncalculated formula yields None; it must not become a number."""
    assert cell_to_text(None) == ""
    assert cell_to_text(float("nan")) == ""
    assert cell_to_text(float("inf")) == ""


def test_a_leading_zero_phone_survives_as_text() -> None:
    assert cell_to_text("01712345678") == "01712345678"


def test_a_spreadsheet_reads_like_the_csv_of_the_same_sheet() -> None:
    content = _workbook(
        [
            ["Customer Mobile", "Name", "Amount"],
            ["01712345678", "Rahim", 1050],
            ["01812345699", "Karim", 1050.5],
        ]
    )

    headers, rows = read_xlsx(content, max_rows=100)

    assert headers == ["Customer Mobile", "Name", "Amount"]
    assert rows[0] == {
        "Customer Mobile": "01712345678",
        "Name": "Rahim",
        "Amount": "1050",
    }
    assert rows[1]["Amount"] == "1050.5"


def test_trailing_blank_rows_are_not_reported_as_failures() -> None:
    """Excel leaves them behind; they are not something a seller did wrong."""
    content = _workbook([["Name", "Amount"], ["Rahim", 1050], [None, None], [None, None]])

    _headers, rows = read_xlsx(content, max_rows=100)

    assert len(rows) == 1


def test_a_sheet_with_no_headings_is_refused_with_a_reason() -> None:
    content = _workbook([[None, None], [None, None]])

    with pytest.raises(ValidationError, match="empty"):
        read_xlsx(content, max_rows=100)


def test_a_sheet_over_the_row_limit_is_refused_rather_than_truncated() -> None:
    """Silently importing the first N rows would lose a seller's data."""
    content = _workbook([["Name"], *[[f"row {index}"] for index in range(30)]])

    with pytest.raises(ValidationError, match="limit"):
        read_xlsx(content, max_rows=10)


def test_only_the_first_sheet_is_read() -> None:
    """Guessing which of three sheets holds the orders imports the wrong data."""
    workbook = Workbook()
    first = workbook.active
    first.append(["Name"])
    first.append(["From sheet one"])
    second = workbook.create_sheet("Other")
    second.append(["Name"])
    second.append(["From sheet two"])
    buffer = io.BytesIO()
    workbook.save(buffer)

    _headers, rows = read_xlsx(buffer.getvalue(), max_rows=100)

    assert [row["Name"] for row in rows] == ["From sheet one"]


# ------------------------------------------------------- upload validation --


def test_a_spreadsheet_is_still_refused_for_a_caller_that_cannot_read_one() -> None:
    """Opt-in, because accepting XLSX is only safe when the reader is safe."""
    content = _workbook([["Name"], ["Rahim"]])

    with pytest.raises(ValidationError, match="Excel"):
        validate_upload(content, filename="orders.xlsx")


def test_a_spreadsheet_is_accepted_when_the_caller_opted_in() -> None:
    content = _workbook([["Name"], ["Rahim"]])

    checked = validate_upload(content, filename="orders.xlsx", allow_spreadsheet=True)

    assert checked.detected is DetectedFormat.XLSX
    # The bytes are returned unread; parsing belongs to the caller.
    assert checked.text == ""


def test_a_zip_that_is_not_a_workbook_is_refused() -> None:
    """ODS and DOCX sniff identically to XLSX; reading either gives nonsense."""
    content = _workbook([["Name"], ["Rahim"]])

    with pytest.raises(ValidationError, match="compressed"):
        validate_upload(content, filename="archive.zip", allow_spreadsheet=True)


# ------------------------------------------------------------ saved mappings --


@pytest.mark.asyncio
async def test_a_mapping_can_be_saved_and_applied_to_a_later_import(db, service) -> None:
    await _shop(db)

    saved = await service.save_mapping(
        name="Weekly orders",
        template=ImportTemplate.ORDERS,
        mapping={"phone": "Customer Mobile", "cod_amount_paisa": "Amount"},
        source_headers=["Customer Mobile", "Amount"],
    )
    batch = ImportBatch(
        template=str(ImportTemplate.ORDERS),
        status=str(ImportStatus.UPLOADED),
        original_filename="orders.csv",
        source_sha256="a" * 64,
        detected_headers=["Customer Mobile", "Amount"],
    )
    db.add(batch)
    await db.flush()

    applied = await service.apply_mapping(batch.id, saved.id)

    assert applied.column_mapping == {
        "phone": "Customer Mobile",
        "cod_amount_paisa": "Amount",
    }
    assert applied.import_status is ImportStatus.MAPPED
    assert saved.use_count == 1
    assert saved.last_used_at is not None


@pytest.mark.asyncio
async def test_saving_the_same_name_updates_rather_than_duplicating(db, service) -> None:
    """Otherwise a seller ends up choosing between two "Daily orders"."""
    await _shop(db)

    first = await service.save_mapping(
        name="Daily orders",
        template=ImportTemplate.ORDERS,
        mapping={"phone": "Mobile"},
    )
    second = await service.save_mapping(
        name="Daily orders",
        template=ImportTemplate.ORDERS,
        mapping={"phone": "Customer Mobile"},
    )

    assert second.id == first.id
    assert second.mapping == {"phone": "Customer Mobile"}
    rows = (await db.execute(sa.select(ImportMapping))).scalars().all()
    assert len(rows) == 1


@pytest.mark.asyncio
async def test_a_mapping_is_only_offered_when_its_columns_are_present(db, service) -> None:
    """Offering one that cannot apply is worse than offering none."""
    await _shop(db)
    await service.save_mapping(
        name="Courier export",
        template=ImportTemplate.ORDERS,
        mapping={"phone": "Customer Mobile", "cod_amount_paisa": "Amount"},
    )

    fits = await service.list_mappings(
        template=ImportTemplate.ORDERS,
        headers=["Customer Mobile", "Amount", "Extra column"],
    )
    does_not = await service.list_mappings(
        template=ImportTemplate.ORDERS, headers=["Phone", "Total"]
    )

    # A superset is fine: a seller who added a column keeps their mapping.
    assert len(fits) == 1
    assert does_not == []


@pytest.mark.asyncio
async def test_a_mapping_for_another_template_cannot_be_applied(db, service) -> None:
    await _shop(db)
    saved = await service.save_mapping(
        name="Products",
        template=ImportTemplate.PRODUCTS,
        mapping={"name": "Item"},
    )
    batch = ImportBatch(
        template=str(ImportTemplate.ORDERS),
        status=str(ImportStatus.UPLOADED),
        original_filename="orders.csv",
        source_sha256="b" * 64,
        detected_headers=["Item"],
    )
    db.add(batch)
    await db.flush()

    with pytest.raises(ValidationError, match="different kind"):
        await service.apply_mapping(batch.id, saved.id)


@pytest.mark.asyncio
async def test_a_mapping_needs_a_name_and_at_least_one_column(db, service) -> None:
    await _shop(db)

    with pytest.raises(ValidationError, match="name"):
        await service.save_mapping(
            name="  ", template=ImportTemplate.ORDERS, mapping={"phone": "Mobile"}
        )
    with pytest.raises(ValidationError, match="at least one"):
        await service.save_mapping(name="Empty", template=ImportTemplate.ORDERS, mapping={})


# -------------------------------------------------------------- error export --


async def _batch_with_failures(db) -> ImportBatch:
    batch = ImportBatch(
        template=str(ImportTemplate.ORDERS),
        status=str(ImportStatus.VALIDATED),
        original_filename="orders.csv",
        source_sha256="c" * 64,
        detected_headers=["Customer Mobile", "Name", "Amount"],
        row_count=3,
    )
    db.add(batch)
    await db.flush()

    db.add_all(
        [
            ImportRow(
                import_id=batch.id,
                row_number=2,
                raw={"Customer Mobile": "not a phone", "Name": "Rahim", "Amount": "1050"},
                status=str(ImportRowStatus.INVALID),
                errors=[{"field": "phone", "message": "That is not a valid mobile number"}],
            ),
            ImportRow(
                import_id=batch.id,
                row_number=3,
                raw={"Customer Mobile": "01712345678", "Name": "Karim", "Amount": "১০৫০"},
                status=str(ImportRowStatus.DUPLICATE),
                errors=[],
            ),
            ImportRow(
                import_id=batch.id,
                row_number=4,
                raw={"Customer Mobile": "01812345699", "Name": "Fine", "Amount": "900"},
                status=str(ImportRowStatus.READY),
                errors=[],
            ),
        ]
    )
    await db.flush()
    return batch


@pytest.mark.asyncio
async def test_the_error_export_contains_only_the_rows_that_failed(db, service) -> None:
    await _shop(db)
    batch = await _batch_with_failures(db)

    body = await service.error_csv(batch.id)

    assert "not a phone" in body
    assert "Karim" in body
    # The row that would be created is not an error.
    assert "Fine" not in body


@pytest.mark.asyncio
async def test_the_error_export_keeps_the_sellers_own_columns(db, service) -> None:
    """So the download is an input — fix, delete two columns, re-upload."""
    await _shop(db)
    batch = await _batch_with_failures(db)

    body = await service.error_csv(batch.id)
    header_line = body.lstrip("﻿").splitlines()[0]

    assert '"Customer Mobile","Name","Amount"' in header_line
    assert "_row_number" in header_line
    assert "_why_it_failed" in header_line


@pytest.mark.asyncio
async def test_the_error_export_says_why_each_row_failed(db, service) -> None:
    await _shop(db)
    batch = await _batch_with_failures(db)

    body = await service.error_csv(batch.id)

    assert "That is not a valid mobile number" in body
    # A duplicate carries no per-field error, so it gets its own reason rather
    # than an empty cell the seller cannot act on.
    assert "Already imported" in body


@pytest.mark.asyncio
async def test_the_error_export_carries_the_row_numbers_from_the_file(db, service) -> None:
    await _shop(db)
    batch = await _batch_with_failures(db)

    body = await service.error_csv(batch.id)

    assert '"2"' in body
    assert '"3"' in body


@pytest.mark.asyncio
async def test_the_error_export_opens_correctly_in_excel(db, service) -> None:
    """A BOM, or Bangla text shows as mojibake and a seller reports it."""
    await _shop(db)
    batch = await _batch_with_failures(db)

    body = await service.error_csv(batch.id)

    assert body.startswith("﻿")
    assert "\r\n" in body
    assert "১০৫০" in body


# ----------------------------------------------------------------- history --


@pytest.mark.asyncio
async def test_history_returns_this_shops_imports_newest_first(db, service) -> None:
    await _shop(db)
    for index in range(3):
        db.add(
            ImportBatch(
                template=str(ImportTemplate.ORDERS),
                status=str(ImportStatus.COMMITTED),
                original_filename=f"orders-{index}.csv",
                source_sha256=f"{index}" * 64,
                created_at=utc_now() + dt.timedelta(minutes=index),
            )
        )
    await db.flush()

    rows = await service.history()

    assert [row.original_filename for row in rows] == [
        "orders-2.csv",
        "orders-1.csv",
        "orders-0.csv",
    ]


@pytest.mark.asyncio
async def test_history_includes_the_failed_ones(db, service) -> None:
    """Those are the imports a seller comes looking for."""
    await _shop(db)
    db.add(
        ImportBatch(
            template=str(ImportTemplate.ORDERS),
            status=str(ImportStatus.FAILED),
            original_filename="broken.csv",
            source_sha256="f" * 64,
        )
    )
    await db.flush()

    rows = await service.history()

    assert [row.status for row in rows] == [str(ImportStatus.FAILED)]


# ------------------------------------------------------------ async commit --


async def _validated_batch(db, *, ready: int) -> ImportBatch:
    batch = ImportBatch(
        template=str(ImportTemplate.ORDERS),
        status=str(ImportStatus.VALIDATED),
        original_filename="big.csv",
        source_sha256="d" * 64,
        row_count=ready,
        ready_count=ready,
    )
    db.add(batch)
    await db.flush()
    return batch


@pytest.mark.asyncio
async def test_a_large_import_is_handed_to_the_worker(db, service) -> None:
    await _shop(db)
    batch = await _validated_batch(db, ready=ASYNC_COMMIT_THRESHOLD_ROWS)

    queued = await service.begin_async_commit(batch.id)

    assert queued.import_status is ImportStatus.COMMITTING
    assert queued.started_at is not None
    events = (
        (
            await db.execute(
                sa.select(OutboxEvent).where(
                    OutboxEvent.topic == str(OutboxTopic.IMPORT_COMMIT_REQUESTED)
                )
            )
        )
        .scalars()
        .all()
    )
    assert len(events) == 1
    assert events[0].payload["import_id"] == str(batch.id)


@pytest.mark.asyncio
async def test_the_job_is_enqueued_through_the_outbox_not_a_second_queue(db, service) -> None:
    """So it exists if and only if the request's transaction commits.

    A direct enqueue would leave a worker committing an import the seller never
    confirmed, every time a request rolled back.
    """
    await _shop(db)
    batch = await _validated_batch(db, ready=ASYNC_COMMIT_THRESHOLD_ROWS)

    await service.begin_async_commit(batch.id)

    event = (
        await db.execute(
            sa.select(OutboxEvent).where(
                OutboxEvent.topic == str(OutboxTopic.IMPORT_COMMIT_REQUESTED)
            )
        )
    ).scalar_one()
    # One job per import, whatever happens upstream.
    assert event.dedupe_key == f"import-commit:{batch.id}"


@pytest.mark.asyncio
async def test_a_second_tap_does_not_enqueue_a_second_job(db, service) -> None:
    await _shop(db)
    batch = await _validated_batch(db, ready=ASYNC_COMMIT_THRESHOLD_ROWS)

    await service.begin_async_commit(batch.id)
    again = await service.begin_async_commit(batch.id)

    assert again.import_status is ImportStatus.COMMITTING
    count = (
        await db.execute(
            sa.select(sa.func.count())
            .select_from(OutboxEvent)
            .where(OutboxEvent.topic == str(OutboxTopic.IMPORT_COMMIT_REQUESTED))
        )
    ).scalar_one()
    assert count == 1


@pytest.mark.asyncio
async def test_an_already_committed_import_is_refused(db, service) -> None:
    await _shop(db)
    batch = await _validated_batch(db, ready=10)
    batch.status = str(ImportStatus.COMMITTED)
    await db.flush()

    with pytest.raises(ConflictError, match="already committed"):
        await service.begin_async_commit(batch.id)


@pytest.mark.asyncio
async def test_a_commit_that_stopped_partway_can_be_resumed(db, service) -> None:
    """A worker killed mid-commit must not strand the batch.

    ``can_commit`` requires VALIDATED, which a COMMITTING batch can never
    return to, so without resumability it would be stuck forever.
    """
    await _shop(db)
    batch = await _validated_batch(db, ready=5)
    batch.status = str(ImportStatus.COMMITTING)
    batch.started_at = utc_now()
    await db.flush()

    assert batch.is_resumable
    assert not batch.can_commit
    # The commit path accepts it, rather than refusing as "not validated".
    committed = await service.commit(batch.id)
    assert committed.import_status is ImportStatus.COMMITTED


@pytest.mark.asyncio
async def test_a_resumed_commit_does_not_recreate_rows(db, service) -> None:
    """The row selection is the idempotency: CREATED rows are never re-taken."""
    await _shop(db)
    batch = await _validated_batch(db, ready=2)
    batch.status = str(ImportStatus.COMMITTING)
    batch.created_count = 1
    await db.flush()

    db.add_all(
        [
            ImportRow(
                import_id=batch.id,
                row_number=2,
                raw={},
                parsed={},
                status=str(ImportRowStatus.CREATED),
                created_entity_id=uuid.uuid4(),
            ),
        ]
    )
    await db.flush()

    committed = await service.commit(batch.id)

    # Nothing left to create, and the earlier row's count is preserved rather
    # than overwritten by this run's zero.
    assert committed.created_count == 1
    remaining = (
        (
            await db.execute(
                sa.select(ImportRow).where(
                    ImportRow.import_id == batch.id,
                    ImportRow.status == str(ImportRowStatus.CREATED),
                )
            )
        )
        .scalars()
        .all()
    )
    assert len(remaining) == 1
